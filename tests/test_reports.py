import pytest
from datetime import datetime, timezone, timedelta, date
import openpyxl
from sqlalchemy.ext.asyncio import AsyncSession

from app.database.models.channel import Channel
from app.database.models.advertisement import Advertisement
from app.services.report_service import ReportService
from app.reports.excel_generator import generate_ads_excel_report


@pytest.mark.asyncio
async def test_report_service_calculations(db_session: AsyncSession):
    # Setup 2 channels
    ch1 = Channel(
        telegram_channel_id=-1001,
        title="Kanal A",
        username="kanal_a",
        url="https://t.me/kanal_a",
        cpm_price=300000.0,
        is_active=True,
    )
    ch2 = Channel(
        telegram_channel_id=-1002,
        title="Kanal B",
        username="kanal_b",
        url="https://t.me/kanal_b",
        cpm_price=None,
        is_active=True,
    )
    db_session.add_all([ch1, ch2])
    await db_session.commit()
    await db_session.refresh(ch1)
    await db_session.refresh(ch2)

    now = datetime.now(timezone.utc)
    # Ad 1 in Channel A (2 days ago, 20,000 views)
    ad1 = Advertisement(
        channel_id=ch1.id,
        message_id=1,
        post_url="https://t.me/kanal_a/1",
        post_text="Reklama 1 @imkonhelpbot",
        posted_at=now - timedelta(days=2),
        views=20000,
        reactions=100,
        forwards=20,
    )
    # Ad 2 in Channel A (4 days ago, 10,000 views)
    ad2 = Advertisement(
        channel_id=ch1.id,
        message_id=2,
        post_url="https://t.me/kanal_a/2",
        post_text="Reklama 2 @imkonhelpbot",
        posted_at=now - timedelta(days=4),
        views=10000,
        reactions=50,
        forwards=10,
    )
    # Ad 3 in Channel B (15 days ago, 5,000 views)
    ad3 = Advertisement(
        channel_id=ch2.id,
        message_id=3,
        post_url="https://t.me/kanal_b/3",
        post_text="Reklama 3 @imkonhelpbot",
        posted_at=now - timedelta(days=15),
        views=5000,
        reactions=25,
        forwards=5,
    )
    # Ad 4 (100 days ago, 8,000 views)
    ad4 = Advertisement(
        channel_id=ch2.id,
        message_id=4,
        post_url="https://t.me/kanal_b/4",
        post_text="Reklama 4 @imkonhelpbot",
        posted_at=now - timedelta(days=100),
        views=8000,
        reactions=40,
        forwards=8,
    )
    db_session.add_all([ad1, ad2, ad3, ad4])
    await db_session.commit()

    service = ReportService(db_session)

    # 1. Test 7 days report (should contain ad1 and ad2)
    rep_7 = await service.get_last_7_days_report()
    assert rep_7["total_ads"] == 2
    assert rep_7["total_views"] == 30000
    assert rep_7["avg_views"] == 15000
    assert rep_7["total_reactions"] == 150
    assert rep_7["total_channels"] == 1
    assert rep_7["top_ad"]["views"] == 20000

    # CPM test for Channel A in rep_7: cost=300000, views=30000 -> CPM = 300000/30000 * 1000 = 10000.0
    ch_stat = rep_7["channel_stats"][0]
    assert ch_stat["cpm"] == 10000.0

    # 2. Test 30 days report (should contain ad1, ad2, ad3)
    rep_30 = await service.get_last_30_days_report()
    assert rep_30["total_ads"] == 3
    assert rep_30["total_views"] == 35000
    assert rep_30["total_channels"] == 2

    # 3. Test 365 days report (should contain all 4 ads)
    rep_365 = await service.get_last_365_days_report()
    assert rep_365["total_ads"] == 4
    assert rep_365["total_views"] == 43000


@pytest.mark.asyncio
async def test_excel_generator_structure_and_styling(db_session: AsyncSession):
    ch = Channel(
        telegram_channel_id=-10099,
        title="Test Kanal",
        username="test_channel",
        url="https://t.me/test_channel",
        is_active=True,
    )
    db_session.add(ch)
    await db_session.commit()
    await db_session.refresh(ch)

    now = datetime.now(timezone.utc)
    ad = Advertisement(
        channel_id=ch.id,
        message_id=99,
        post_url="https://t.me/test_channel/99",
        post_text="=SUM(A1:A10) Chegirma @imkonhelpbot",
        posted_at=now,
        media_type="photo",
        views=25000,
        reactions=120,
        forwards=30,
    )
    db_session.add(ad)
    await db_session.commit()

    service = ReportService(db_session)
    report_data = await service.get_last_7_days_report()

    # Generate Excel buffer
    buffer = generate_ads_excel_report(report_data)
    assert buffer is not None

    # Load back using openpyxl to verify sheets & structure
    wb = openpyxl.load_workbook(buffer)
    sheet_names = wb.sheetnames

    assert "REKLAMALAR" in sheet_names
    assert "KANALLAR STATISTIKASI" in sheet_names
    assert "UMUMIY HISOBOT" in sheet_names
    assert "KUNLIK STATISTIKA" in sheet_names

    ws_ads = wb["REKLAMALAR"]
    assert ws_ads.cell(row=1, column=1).value == "№"
    assert ws_ads.cell(row=2, column=2).value == "Test Kanal"
    assert ws_ads.cell(row=2, column=7).value == 25000
    assert ws_ads.freeze_panes == "A2"

    ws_summary = wb["UMUMIY HISOBOT"]
    assert "IMKON" in str(ws_summary["B2"].value)

    ws_daily = wb["KUNLIK STATISTIKA"]
    assert ws_daily.cell(row=1, column=1).value == "Sana"

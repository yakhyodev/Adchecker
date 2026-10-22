"""Package integrity and data preservation are checked independently of Excel."""

from io import BytesIO
from datetime import date, datetime
from threading import Event
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from app.reports.post_excel import HEADERS, _row, generate_post_excel
from app.core.timezone import TASHKENT_TZ
from app.reports.post_range_excel import FORWARD, SOURCE, MISSING, generate_post_range_excel
from app.reports.xlsx_validation import validate_xlsx
from app.services.post_range import RangeStore
from app.services.post_report import PostReport, metrics
from tests.test_post_report import NOW, channel, post

CASES = ["empty", "one", "many", "partial", "missing", "special"]


def make_report(case):
    source = metrics(post(), channel())
    count = 0 if case == "empty" else 3 if case == "many" else 1
    forwards = [metrics(post(10 + n, 101 + n, views=n * 100, shares=n),
                        channel(101 + n, f"destination_{n}")) for n in range(count)]
    if case == "missing":
        source.views = source.reactions = source.replies = source.members = None
        forwards[0].views = forwards[0].shares = forwards[0].members = None
    if case == "special":
        source.title = "IMKON & <savdo> \"O'zbek\" ⚡"
        forwards[0].title = '=HYPERLINK("https://bad.example") ' + "Uzun nom ⚡ <>&'\" " * 30 + "\x01\ufffe"
    return PostReport(source=source, checked_at=NOW, public_shares=count, private_shares=228-count,
                      forwards=forwards, forwards_complete=case != "partial")


def assert_package(path):
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                ET.fromstring(archive.read(name))
    if hasattr(path, "seek"):
        path.seek(0)
    wb = load_workbook(path)
    for ws in wb:
        for table in ws.tables.values():
            left, top, right, bottom = range_boundaries(table.ref)
            assert bottom > top
            assert [column.name for column in table.tableColumns] == [ws.cell(top, col).value for col in range(left, right + 1)]
            assert len(table.tableColumns) == right - left + 1
    return wb


@pytest.mark.parametrize("case", CASES)
def test_single_post_xlsx_integrity_and_all_existing_fields(case):
    report = make_report(case)
    buffer = generate_post_excel(report)
    wb = assert_package(buffer)
    source, forwards = wb.worksheets
    assert [cell.value for cell in source[1]][:10] == HEADERS
    assert [cell.value for cell in source[2]][:10] == _row(report.source)
    assert source["K2"].value == report.public_shares
    assert source["L2"].value == report.private_shares
    assert len(source.tables) == 1
    assert len(forwards.tables) == bool(report.forwards)
    assert forwards.max_row == len(report.forwards) + 1
    for number, data in enumerate(report.forwards, 2):
        assert [cell.value for cell in forwards[number]] == _row(data)
        assert forwards.cell(number, 5).hyperlink.target == data.post_url
    assert source["E2"].hyperlink.target == report.source.post_url
    assert source["F2"].number_format == "dd.mm.yyyy hh:mm"
    if report.source.views is not None:
        assert source["G2"].number_format == "#,##0"
    assert source.freeze_panes == "C2"
    assert all(cell.data_type != "f" for ws in wb for row in ws for cell in row)
    assert ("qisman" in forwards.title) == (case == "partial")
    wb.close()


@pytest.mark.parametrize("case", CASES)
def test_range_xlsx_integrity_and_requested_columns(case, tmp_path):
    report = make_report(case)
    store = RangeStore(tmp_path / "spool.sqlite3")
    try:
        for row in report.forwards:
            store.sink(report.source.message_id).save(row)
        store.save_post(report, NOW.astimezone(TASHKENT_TZ).date())
        target = tmp_path / "range.xlsx"
        generate_post_range_excel(store, target, Event())
        wb = assert_package(target)
        assert wb.sheetnames == ["Umumiy xulosa", "Postlar", "Forwardlar"]
        posts, forwards = wb["Postlar"], wb["Forwardlar"]
        assert [cell.value for cell in posts[1]] == SOURCE
        assert [cell.value for cell in forwards[1]] == FORWARD
        assert posts["A2"].value == 1
        assert posts["C2"].hyperlink.target == report.source.post_url
        assert posts["D2"].value == (MISSING if report.source.views is None else report.source.views)
        assert posts["F2"].value == report.public_shares
        assert posts["G2"].value == report.private_shares
        assert forwards.max_row == 1 + len(report.forwards)
        expected_rows = sorted(report.forwards, key=lambda item: (item.published_at, item.telegram_id, item.message_id))
        for row, expected in enumerate(expected_rows, 2):
            assert forwards.cell(row, 2).value == 1
            assert forwards.cell(row, 5).hyperlink.target == expected.post_url
            assert forwards.cell(row, 7).value == (MISSING if expected.views is None else expected.views)
        assert not forwards.tables
        summary = dict(wb["Umumiy xulosa"].iter_rows(min_row=2, values_only=True))
        assert summary["Holat"] == ("Qisman" if case == "partial" else "To'liq")
        if case == "missing":
            assert summary["Asosiy postlarning umumiy ko'rishlari"] == MISSING
            assert summary["Forwardlarning umumiy ko'rishlari"] == MISSING
        assert all(cell.data_type != "f" for ws in wb for row in ws for cell in row)
        wb.close()
    finally:
        store.close()


def test_validator_rejects_header_only_tables_even_if_openpyxl_accepts_them():
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table
    wb = Workbook()
    wb.active.append(["Sarlavha"])
    wb.active.add_table(Table(displayName="Empty", ref="A1:A1"))
    buffer = BytesIO()
    wb.save(buffer)
    with pytest.raises(ValueError, match="diapazoni"):
        validate_xlsx(buffer)


def test_day_bands_group_every_post_in_a_day_and_restart_blue_on_third_day(tmp_path):
    store = RangeStore(tmp_path / "spool.sqlite3")
    try:
        for identifier, day in enumerate([1, 1, 2, 3, 3], 100):
            report = make_report("empty")
            report.source.message_id = identifier
            report.source.published_at = datetime(2026, 10, day, 12, tzinfo=TASHKENT_TZ)
            store.save_post(report, date(2026, 10, day))
        path = tmp_path / "days.xlsx"
        generate_post_range_excel(store, path, Event())
        wb = assert_package(path)
        assert [wb["Postlar"].cell(row, 1).fill.fgColor.rgb for row in range(2, 7)] == [
            "00EAF3F8", "00EAF3F8", "00FFFFFF", "00EAF3F8", "00EAF3F8"]
        assert [wb["Postlar"].cell(row, 1).value for row in range(2, 7)] == [1, 2, 3, 4, 5]
        wb.close()
    finally:
        store.close()

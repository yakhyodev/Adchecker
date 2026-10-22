import asyncio
import logging
import tracemalloc
from copy import deepcopy
from datetime import date, datetime, timezone
from threading import Event
from types import SimpleNamespace as NS
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook
from telethon.errors import ChatAdminRequiredError, FloodWaitError, FloodPremiumWaitError
from telethon.tl.functions.channels import GetFullChannelRequest
from telethon.tl.functions.stats import GetMessagePublicForwardsRequest
from telethon.tl.types import PublicForwardMessage

from app.core.timezone import TASHKENT_TZ
from app.reports.post_range_excel import generate_post_range_excel
from app.services.post_range import PostRange, PostRangeService, RangeStore, SequentialPostService, decode
from app.services.post_report import MISSING
from app.services.post_report import PostReport, metrics
from tests.test_post_report import FakeClient, channel, post


@pytest.fixture(autouse=True)
def freeze_now(monkeypatch):
    monkeypatch.setattr("app.services.post_range.get_tashkent_now", lambda: datetime(2026, 10, 22, 12, tzinfo=TASHKENT_TZ))


def dated(identifier, day, hour=0, minute=0, second=0):
    result = post(identifier)
    result.date = datetime(2026, 10, day, hour, minute, second, tzinfo=TASHKENT_TZ).astimezone(timezone.utc)
    return result


class RangeClient(FakeClient):
    def __init__(self, messages):
        super().__init__()
        self.history = messages
        self.history_calls = []
        self.stats_calls = []
        self.placements = {}
        self.fail_stats_once = False
        self.fail_history_once = False
        self.active = self.max_active = 0
        self.rpc.side_effect = self.range_respond
        self.get_messages.side_effect = lambda entity, ids: deepcopy(next(item for item in self.history if item.id == ids))

    async def range_respond(self, request):
        self.active += 1
        self.max_active = max(self.max_active, self.active)
        try:
            await asyncio.sleep(0)
            if isinstance(request, GetFullChannelRequest):
                return self.respond(request)
            if isinstance(request, GetMessagePublicForwardsRequest):
                self.stats_calls.append((request.msg_id, request.offset))
                if self.fail_stats_once:
                    self.fail_stats_once = False
                    raise FloodWaitError(None, capture=32)
                return self.placements.get(request.msg_id, self.pages[""])
            raise AssertionError(type(request).__name__)
        finally:
            self.active -= 1

    def iter_messages(self, entity, **kwargs):
        self.history_calls.append(kwargs)
        assert kwargs["reverse"] and kwargs["limit"] is None

        async def history():
            for message in self.history:
                if message.id <= kwargs["min_id"]:
                    continue
                if kwargs["offset_date"] and message.date <= kwargs["offset_date"]:
                    continue
                if self.fail_history_once and message.id == 11:
                    self.fail_history_once = False
                    raise FloodWaitError(None, capture=4)
                yield message
        return history()


@pytest.fixture
def store(tmp_path):
    result = RangeStore(tmp_path / "spool.sqlite3")
    yield result
    result.close()


def service(client, sleep=None):
    manager = NS(client=client, probe_authorization=AsyncMock(return_value=True))
    return PostRangeService(manager, pace=0, sleep=sleep or AsyncMock())


@pytest.mark.asyncio
async def test_oldest_first_local_midnight_inclusive_and_no_parallel_requests(store):
    client = RangeClient([dated(9, 9, 23, 59, 59), dated(10, 10), dated(11, 10, 23, 59, 59),
                          dated(12, 11), dated(13, 11, 12), dated(14, 12)])
    await service(client).collect(PostRange(date(2026, 10, 10), date(2026, 10, 11)), store, AsyncMock())
    assert [row[0] for row in store.db.execute("SELECT message_id FROM posts ORDER BY seq")] == [10, 11, 12, 13]
    assert client.stats_calls == [(10, ""), (11, ""), (12, ""), (13, "")]
    assert client.max_active == 1
    assert [call.kwargs["ids"] for call in client.get_messages.await_args_list] == [10, 11, 12, 13]
    assert len(client.history_calls) == 2


@pytest.mark.asyncio
async def test_public_stats_waits_and_retries_same_request_before_next_post(store):
    client = RangeClient([dated(10, 10), dated(11, 10, 12)])
    client.fail_stats_once = True
    sleep, progress = AsyncMock(), AsyncMock()
    await service(client, sleep).collect(PostRange(date(2026, 10, 10), date(2026, 10, 10)), store, progress)
    assert client.stats_calls == [(10, ""), (10, ""), (11, "")]
    assert [call.args[0] for call in sleep.await_args_list] == [30, 3]
    assert any(call.args[0] == "wait" for call in progress.await_args_list)
    assert any(call.args[0] == "resume" for call in progress.await_args_list)
    assert store.post_count == 2 and store.partial_count == 0


@pytest.mark.asyncio
async def test_history_generator_reopens_at_last_completed_id_after_floodwait(store):
    client = RangeClient([dated(10, 10), dated(11, 10, 12), dated(12, 10, 16)])
    client.fail_history_once = True
    await service(client).collect(PostRange(date(2026, 10, 10), date(2026, 10, 10)), store, AsyncMock())
    assert [call["min_id"] for call in client.history_calls] == [0, 10]
    assert client.stats_calls == [(10, ""), (11, ""), (12, "")]
    assert store.post_count == 3


@pytest.mark.asyncio
async def test_wait_can_be_cancelled_without_further_requests(store):
    client = RangeClient([dated(10, 10)])
    client.fail_stats_once = True
    waiting = asyncio.Event()

    async def blocked(seconds):
        waiting.set()
        await asyncio.Future()

    task = asyncio.create_task(service(client, blocked).collect(
        PostRange(date(2026, 10, 10), date(2026, 10, 10)), store, AsyncMock()))
    await asyncio.wait_for(waiting.wait(), 2)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert client.stats_calls == [(10, "")] and store.post_count == 0


@pytest.mark.asyncio
async def test_date_post_column_blocks_numeric_id_order_and_multiple_forwards(store, tmp_path):
    client = RangeClient([dated(10, 10), dated(11, 11)])
    first = post(5, 101, 12, 3)
    second = post(6, 101, 19, 4)
    other = post(7, 201, 28, 5)
    client.placements = {
        10: NS(count=3, forwards=[PublicForwardMessage(second), PublicForwardMessage(other), PublicForwardMessage(first)],
               chats=[channel(101, "AAA"), channel(201, "ZZZ")], next_offset=None),
        11: NS(count=1, forwards=[PublicForwardMessage(other)], chats=[channel(201, "ZZZ")], next_offset=None),
    }
    await service(client).collect(PostRange(date(2026, 10, 10), date(2026, 10, 11)), store, AsyncMock())
    path = tmp_path / "range.xlsx"
    generate_post_range_excel(store, path, Event())
    wb = load_workbook(path)
    source, forwards = wb.worksheets
    assert source.title == "Asosiy postlar" and forwards.title == "Public forwardlar"
    assert "2026-10-10 · Post #10" in source.cell(1, 5).value
    assert "2026-10-11 · Post #11" in source.cell(1, 15).value
    assert [forwards.cell(row, 1).value for row in (2, 3)] == [-1000000000201, -1000000000101]
    assert forwards.cell(2, 2).value == "Kanal 201"
    assert forwards.cell(3, 5).value == "https://t.me/AAA/5"
    assert forwards.cell(3, 12).value == "https://t.me/AAA/6"
    assert "2026-10-11 · Post #11" in forwards.cell(1, 19).value
    assert forwards.cell(3, 19).value == MISSING
    assert source.cell(2, 6).value == datetime(2026, 10, 10)
    assert forwards.auto_filter.ref.endswith("3")
    assert not forwards._hyperlinks and not forwards._charts
    assert all(cell.data_type != "f" for ws in wb for row in ws for cell in row)
    assert not list(store.db.execute("SELECT 1 FROM posts WHERE complete=0"))


@pytest.mark.asyncio
async def test_xlsx_column_limit_is_split_without_losing_post_blocks(store, tmp_path):
    client = RangeClient([dated(10, 10), dated(11, 10, 12), dated(12, 11)])
    await service(client).collect(PostRange(date(2026, 10, 10), date(2026, 10, 11)), store, AsyncMock())
    path = tmp_path / "split.xlsx"
    generate_post_range_excel(store, path, Event(), max_columns=24)
    wb = load_workbook(path)
    assert len(wb.worksheets) == 4
    assert all(ws.max_column <= 24 for ws in wb)
    assert "Post #12" in wb["Asosiy postlar 2"].cell(1, 5).value


@pytest.mark.asyncio
async def test_large_single_post_streams_forwards_to_disk_not_report_list(store):
    client = FakeClient()
    pages = 25

    def respond(request):
        if isinstance(request, GetFullChannelRequest):
            return client.respond(request)
        index = int(request.offset or 0)
        chats = [channel(1000 + n, f"dest{n}") for n in range(index * 100, (index + 1) * 100)]
        return NS(count=pages * 100, forwards=[PublicForwardMessage(post(n + 1, chat.id)) for n, chat in enumerate(chats)],
                  chats=chats, next_offset=str(index + 1) if index + 1 < pages else None)

    client.rpc.side_effect = respond
    api = SequentialPostService(NS(client=client), AsyncMock(), pace=0, sleep=AsyncMock())
    tracemalloc.start()
    try:
        report = await api.collect_message(client.source, post(), store.sink(5644))
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert report.forwards == [] and report.forward_count == pages * 100
    assert store.db.execute("SELECT COUNT(*) FROM forwards").fetchone()[0] == pages * 100
    assert peak < 12 * 1024 * 1024


@pytest.mark.asyncio
async def test_unavailable_stats_are_partial_information_not_retry_forever(store, caplog):
    client = RangeClient([dated(10, 10)])
    client.placements[10] = ChatAdminRequiredError(None)
    original = client.range_respond

    async def respond(request):
        if isinstance(request, GetMessagePublicForwardsRequest):
            raise ChatAdminRequiredError(None)
        return await original(request)

    client.rpc.side_effect = respond
    with caplog.at_level(logging.INFO, logger="imkon_ads_monitor"):
        await service(client).collect(PostRange(date(2026, 10, 10), date(2026, 10, 10)), store, AsyncMock())
    assert store.post_count == store.partial_count == 1
    assert all(record.levelno < logging.WARNING for record in caplog.records)


@pytest.mark.parametrize("start,end", [(date(2026, 10, 12), date(2026, 10, 10)),
                                       (date(2026, 10, 10), date(2026, 10, 23))])
def test_invalid_ranges(start, end):
    with pytest.raises(ValueError):
        PostRange.create(start, end)


def test_streaming_export_memory_is_bounded_for_many_destinations(store, tmp_path):
    source = PostReport(source=metrics(post(), channel()), checked_at=datetime(2026, 10, 21, tzinfo=timezone.utc),
                        forwards_complete=True)
    sink = store.sink(5644)
    for number in range(2500):
        sink.save(metrics(post(number + 1, 1000 + number), channel(1000 + number, f"dest{number}")))
    store.save_post(source, date(2026, 10, 21))
    path = tmp_path / "large.xlsx"
    tracemalloc.start()
    try:
        generate_post_range_excel(store, path, Event())
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    print(f"2500-destination XLSX export peak Python allocations: {peak / 1024**2:.2f} MiB")
    assert peak < 16 * 1024 * 1024
    wb = load_workbook(path, read_only=True)
    try:
        assert sum(1 for _ in wb["Public forwardlar"].rows) == 2501
    finally:
        wb.close()


@pytest.mark.asyncio
async def test_other_telegram_rate_limit_is_waited_and_retried():
    attempts = 0

    async def request():
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise FloodPremiumWaitError(None, capture=8)
        return "ok"

    sleep = AsyncMock()
    api = SequentialPostService(NS(client=None), AsyncMock(), pace=0, sleep=sleep)
    assert await api._optional(request, "statistika") == "ok"
    sleep.assert_awaited_once_with(9)
    assert attempts == 2


@pytest.mark.asyncio
async def test_source_statistics_are_refreshed_when_each_posts_turn_arrives(store):
    client = RangeClient([dated(10, 10), dated(11, 10, 12)])

    async def fresh(entity, ids):
        result = deepcopy(next(item for item in client.history if item.id == ids))
        result.views = 900000 + ids
        return result

    client.get_messages.side_effect = fresh
    await service(client).collect(PostRange(date(2026, 10, 10), date(2026, 10, 10)), store, AsyncMock())
    sources = [decode(row[0]) for row in store.db.execute("SELECT source FROM posts ORDER BY seq")]
    assert [item.views for item in sources] == [900010, 900011]
    assert client.history[0].views == 700300


@pytest.mark.asyncio
async def test_deleted_post_keeps_only_metadata_and_does_not_stop_remaining_posts(store):
    client = RangeClient([dated(10, 10), dated(11, 10, 12)])
    client.get_messages.side_effect = lambda entity, ids: None if ids == 10 else deepcopy(client.history[1])
    await service(client).collect(PostRange(date(2026, 10, 10), date(2026, 10, 10)), store, AsyncMock())
    first = decode(store.db.execute("SELECT source FROM posts WHERE message_id=10").fetchone()[0])
    assert first.views is first.shares is first.reactions is None
    assert first.post_url.endswith("/10")
    assert store.post_count == 2 and store.partial_count == 1
    assert client.stats_calls == [(11, "")]


@pytest.mark.asyncio
async def test_authorization_rate_limit_is_retried_instead_of_becoming_logged_out(store):
    client = RangeClient([dated(10, 10)])
    manager = NS(client=client, probe_authorization=AsyncMock(side_effect=[FloodWaitError(None, capture=3), True]))
    sleep = AsyncMock()
    await PostRangeService(manager, pace=0, sleep=sleep).collect(
        PostRange(date(2026, 10, 10), date(2026, 10, 10)), store, AsyncMock())
    assert manager.probe_authorization.await_count == 2 and store.post_count == 1
    sleep.assert_awaited_once_with(4)


def test_cancelled_export_removes_xml_from_job_directory(store, tmp_path, monkeypatch):
    from app.reports import post_range_excel
    source = PostReport(source=metrics(post(), channel()), checked_at=datetime(2026, 10, 21, tzinfo=timezone.utc),
                        forwards_complete=True)
    store.save_post(source, date(2026, 10, 21))
    stop = Event()
    original = post_range_excel._cells

    def cells(ws, values, header=False):
        yield from original(ws, values, header)
        if not header:
            stop.set()

    monkeypatch.setattr(post_range_excel, "_cells", cells)
    with pytest.raises(post_range_excel.RangeExportCancelled):
        generate_post_range_excel(store, tmp_path / "cancelled.xlsx", stop)
    assert not list(tmp_path.glob("*.xml")) and not (tmp_path / "cancelled.xlsx").exists()

import asyncio
from datetime import date
from threading import Event
from unittest.mock import AsyncMock

import pytest
from aiogram.exceptions import TelegramNetworkError, TelegramRetryAfter
from aiogram.methods import SendDocument
from openpyxl import load_workbook

from app.bot.handlers import post_range, post_report
from app.bot.operations import cancel_operation, running
from app.services.post_range import PostRange
from app.services.post_range import RangeStore
from app.services.post_report import PostReport, PostReportUnavailable, metrics
from tests.test_post_report import NOW, channel, post
from tests.test_post_report_flow import state_and_message


@pytest.mark.asyncio
async def test_cancel_closes_disk_spool_and_releases_shared_slot(monkeypatch, tmp_path):
    monkeypatch.setattr(post_range, "WORK_ROOT", tmp_path)
    state, message = state_and_message()
    entered = asyncio.Event()

    async def blocked(selected, store, progress):
        store.save_post(PostReport(source=metrics(post(), channel()), checked_at=NOW), selected.start)
        entered.set()
        await asyncio.Future()

    monkeypatch.setattr(post_range.service, "collect", blocked)
    task = asyncio.create_task(post_range.run_range(message, state, PostRange(date(2026, 10, 21), date(2026, 10, 21))))
    await asyncio.wait_for(entered.wait(), 2)
    assert list(tmp_path.iterdir())
    await cancel_operation(777, 777)
    await asyncio.wait_for(task, 2)
    assert not list(tmp_path.iterdir())
    assert not post_report.report_slots.locked() and (777, 777) not in running
    assert await state.get_state() is None
    message.answer_document.assert_not_awaited()


@pytest.mark.asyncio
async def test_single_post_report_waits_behind_batch_worker(monkeypatch, tmp_path):
    monkeypatch.setattr(post_range, "WORK_ROOT", tmp_path)
    first_state, first_message = state_and_message(777)
    second_state, second_message = state_and_message(888)
    entered = asyncio.Event()

    async def blocked(selected, store, progress):
        entered.set()
        await asyncio.Future()

    monkeypatch.setattr(post_range.service, "collect", blocked)
    single = AsyncMock()
    monkeypatch.setattr(post_report.service, "collect", single)
    first = asyncio.create_task(post_range.run_range(first_message, first_state, PostRange(date(2026, 10, 21), date(2026, 10, 21))))
    second = None
    try:
        await asyncio.wait_for(entered.wait(), 2)
        second = asyncio.create_task(post_report.run_report(second_message, second_state, "https://t.me/imkonsavdomarkazi/5644"))
        for _ in range(20):
            if second_message.answer.await_count:
                break
            await asyncio.sleep(0)
        single.assert_not_awaited()
        await cancel_operation(888, 888)
        assert post_report.report_slots.locked()
    finally:
        await cancel_operation(777, 777)
        await asyncio.gather(*(task for task in (first, second) if task is not None), return_exceptions=True)
    assert not post_report.report_slots.locked()


@pytest.mark.asyncio
async def test_cancelled_export_joins_writer_before_database_cleanup(monkeypatch, tmp_path):
    entered, finished = Event(), Event()

    def writer(store, path, stop):
        entered.set()
        stop.wait(2)
        finished.set()

    monkeypatch.setattr(post_range, "generate_post_range_excel", writer)
    task = asyncio.create_task(post_range.export_range(object(), tmp_path / "unused.xlsx"))
    assert await asyncio.to_thread(entered.wait, 2)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert finished.is_set()


@pytest.mark.asyncio
async def test_network_and_telegram_limits_retry_existing_file_without_recollecting(tmp_path):
    _, message = state_and_message()
    path = tmp_path / "ready.xlsx"
    path.write_bytes(b"already generated")
    request = SendDocument(chat_id=777, document="test")
    message.answer_document.side_effect = [TelegramNetworkError(method=request, message="offline"),
        TelegramRetryAfter(method=request, message="wait", retry_after=12), None]
    sleep, progress = AsyncMock(), AsyncMock()
    await post_range.send_file(message, path, progress, sleep=sleep)
    assert message.answer_document.await_count == 3
    assert [call.args[0] for call in sleep.await_args_list] == [5, 13]
    assert path.read_bytes() == b"already generated"


@pytest.mark.asyncio
async def test_file_delivery_wait_can_be_cancelled(tmp_path):
    _, message = state_and_message()
    request = SendDocument(chat_id=777, document="test")
    message.answer_document.side_effect = TelegramNetworkError(method=request, message="offline")
    waiting = asyncio.Event()

    async def sleep(seconds):
        waiting.set()
        await asyncio.Future()

    task = asyncio.create_task(post_range.send_file(message, tmp_path / "ready.xlsx", AsyncMock(), sleep=sleep))
    await asyncio.wait_for(waiting.wait(), 2)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert message.answer_document.await_count == 1


def populate(store, destinations=60):
    store.save_post(PostReport(source=metrics(post(), channel()), checked_at=NOW, forwards_complete=True), date(2026, 10, 21))
    sink = store.sink(5644)
    for number in range(destinations):
        row = metrics(post(number + 1, 1000 + number, views=number * 7321, shares=number),
                      channel(1000 + number, f"destination_{number}"))
        sink.save(row)
    store.db.commit()


@pytest.mark.asyncio
async def test_oversized_workbook_is_split_from_spool_with_all_ids_preserved(tmp_path):
    store = RangeStore(tmp_path / "spool.sqlite3")
    try:
        populate(store)
        big, half = tmp_path / "full.xlsx", tmp_path / "half.xlsx"
        await post_range.export_range(store, big)
        await post_range.export_range(store, half, forward_offset=0, forward_limit=30)
        limit = (big.stat().st_size + half.stat().st_size) // 2
        files = await post_range.export_files(store, tmp_path / "report.xlsx", max_bytes=limit)
        assert len(files) > 1 and all(path.stat().st_size <= limit for path in files)
        actual = []
        for path in files:
            wb = load_workbook(path, read_only=True)
            try:
                actual.extend(row[4] for row in wb["Forwardlar"].iter_rows(min_row=2, values_only=True))
            finally:
                wb.close()
        from app.reports.post_range_excel import forward_rows
        from app.services.post_range import decode
        expected = [decode(row[2]).post_url for row in forward_rows(store)]
        assert actual == expected
        assert not list(tmp_path.glob("*.xml"))
    finally:
        store.close()


@pytest.mark.asyncio
async def test_excel_row_limit_splits_before_writing_invalid_rows(tmp_path):
    store = RangeStore(tmp_path / "spool.sqlite3")
    try:
        populate(store, destinations=3)
        files = await post_range.export_files(store, tmp_path / "report.xlsx", max_rows=1)
        assert len(files) == 3
        for path in files:
            wb = load_workbook(path)
            assert wb["Forwardlar"].max_row == 2
            wb.close()
    finally:
        store.close()


@pytest.mark.asyncio
async def test_splitting_multiple_sources_keeps_every_forward_and_its_source_number(tmp_path):
    store = RangeStore(tmp_path / "spool.sqlite3")
    try:
        for number in range(3):
            source = metrics(post(100 + number), channel())
            store.save_post(PostReport(source=source, checked_at=NOW, forwards_complete=True), date(2026, 10, 21))
            for copy in range(3):
                store.sink(source.message_id).save(metrics(post(1000 + number * 10 + copy, 101), channel(101, "destination")))
        store.db.commit()
        files = await post_range.export_files(store, tmp_path / "report.xlsx", max_rows=2)
        found = []
        for path in files:
            wb = load_workbook(path)
            sources = {row[0]: row[2] for row in wb["Postlar"].iter_rows(min_row=2, values_only=True)}
            assert wb["Postlar"].max_row <= 3 and wb["Forwardlar"].max_row <= 3
            for row in wb["Forwardlar"].iter_rows(min_row=2, values_only=True):
                assert sources[row[1]].endswith(f"/{99 + row[1]}")
                found.append(row[4])
            wb.close()
        assert len(found) == len(set(found)) == 9
    finally:
        store.close()


@pytest.mark.asyncio
async def test_late_history_failure_still_delivers_completed_posts_as_partial(monkeypatch, tmp_path):
    monkeypatch.setattr(post_range, "WORK_ROOT", tmp_path)
    state, message = state_and_message()

    async def interrupted(selected, store, progress):
        store.save_post(PostReport(source=metrics(post(), channel()), checked_at=NOW, forwards_complete=True), selected.start)
        raise PostReportUnavailable("Kanal tarixiga kirish imkoni yo'q.")

    async def uploaded(document):
        assert "qisman" in document.filename
        wb = load_workbook(document.path)
        assert wb.sheetnames == ["Postlar", "Forwardlar"]
        wb.close()

    monkeypatch.setattr(post_range.service, "collect", interrupted)
    message.answer_document.side_effect = uploaded
    await post_range.run_range(message, state, PostRange(date(2026, 10, 21), date(2026, 10, 22)))
    message.answer_document.assert_awaited_once()
    assert any("Qisman hisobot" in call.args[0] for call in message.answer.return_value.edit_text.await_args_list)
    assert not list(tmp_path.iterdir()) and await state.get_state() is None

from datetime import date, datetime, timezone
from types import SimpleNamespace as NS
from unittest.mock import AsyncMock, MagicMock

import openpyxl
import pytest
from sqlalchemy import func, select
from telethon.errors import ChannelPrivateError, FloodWaitError

from app.core.timezone import TASHKENT_TZ
from app.database.models.advertisement import Advertisement
from app.database.models.channel import Channel
from app.database.models.view_snapshot import ViewSnapshot
from app.reports.live_excel import generate_live_excel
from app.services.channel_service import ChannelService
from app.services.live_search import (
    ChannelSearchResult, LiveAd, LiveSearchService, SearchChannel, SearchReport,
    SearchRequest, SearchUnavailable, load_search_channels, parse_keywords,
)
from app.telegram_client.client import ChannelNotFoundError, TelegramFloodWaitException


@pytest.fixture(autouse=True)
def freeze_query_time(monkeypatch):
    monkeypatch.setattr("app.services.live_search.get_tashkent_now", lambda: datetime(2026, 10, 22, 12, tzinfo=TASHKENT_TZ))


async def stream(items):
    for item in items:
        if isinstance(item, Exception):
            raise item
        yield item


def channel(number=111):
    return SearchChannel(number, int(f"-100{number}"), "Old title", "https://t.me/old_name", "old_name")


def post(number, text="@otherbot", at=None, **kwargs):
    return NS(id=number, message=text, date=at or datetime(2026, 10, 6, 10, tzinfo=timezone.utc),
              entities=kwargs.get("entities"), grouped_id=kwargs.get("album"),
              views=kwargs.get("views", 60200), forwards=kwargs.get("forwards", 27))


def manager(messages, peers=None):
    client = MagicMock()
    client.get_entity = AsyncMock(side_effect=lambda marked: NS(
        id=int(str(marked).removeprefix("-100")), username="new_name", title="New title", broadcast=True))
    def iterate(peer, **kw):
        rows = messages.get(peer.id, [])
        query = kw.get("search")
        if query:
            # Model a text index that does not return hidden text-link URLs.
            rows = [msg for msg in rows if isinstance(msg, Exception) or query.casefold() in msg.message.casefold()]
        limit = kw.get("limit")
        return stream(rows[:limit] if limit is not None else rows)
    client.iter_messages.side_effect = iterate
    client.iter_dialogs.side_effect = lambda **kw: stream(peers or [])
    return NS(client=client, is_authorized=AsyncMock(return_value=True))


def custom_request():
    return SearchRequest.create("@otherbot", "custom", date(2026, 10, 1), date(2026, 10, 21))


def test_keywords_and_calendar_month(monkeypatch):
    monkeypatch.setattr("app.services.live_search.get_tashkent_now", lambda: datetime(2026, 10, 22, 12, tzinfo=TASHKENT_TZ))
    req = SearchRequest.create("@OtherBot, IMKON savdo\n@otherbot", "month")
    assert req.keywords == ("@otherbot", "imkon savdo")
    assert req.start == datetime(2026, 9, 30, 19, tzinfo=timezone.utc)
    assert req.end == datetime(2026, 10, 22, 7, tzinfo=timezone.utc)
    with pytest.raises(ValueError):
        parse_keywords(" , \n")
    with pytest.raises(ValueError):
        SearchRequest.create("a", "custom", date(2026, 10, 22), date(2026, 10, 1))
    with pytest.raises(ValueError):
        SearchRequest.create("a", "custom", date(2026, 10, 1), date(2026, 10, 23))


@pytest.mark.asyncio
async def test_latest_uses_saved_id_after_rename_and_stops_at_latest_match():
    # The history is newest-first; the latest POST is unrelated, but the next is a match.
    mgr = manager({111: [post(4, "news"), post(3, "@OTHERBOT"), post(2)]})
    report = await LiveSearchService(mgr).search([channel()], SearchRequest.create("@otherbot", "latest"))
    mgr.client.get_entity.assert_awaited_once_with(-100111)
    mgr.client.iter_dialogs.assert_not_called()
    assert [ad.message_id for ad in report.channels[0].ads] == [3]
    assert report.channels[0].ads[0].url == "https://t.me/new_name/3"
    assert report.channels[0].channel.title == "New title"


@pytest.mark.asyncio
async def test_empty_entity_cache_resolves_only_saved_ids_from_dialog_metadata():
    mgr = manager({111: [post(1)]}, peers=[
        NS(entity=NS(id=999, title="Unregistered", username="unregistered", broadcast=True)),
        NS(entity=NS(id=111, title="Renamed", username=None, broadcast=True)),
        NS(entity=NS(id=111, title="User with same raw ID", username=None, broadcast=False)),
    ])
    mgr.client.get_entity.side_effect = ValueError("Missing access hash")
    report = await LiveSearchService(mgr).search([channel()], custom_request())
    assert report.channels[0].ads[0].url == "https://t.me/c/111/1"
    assert mgr.client.iter_messages.call_count == 2
    assert mgr.client.iter_messages.call_args.args[0].id == 111


@pytest.mark.asyncio
async def test_date_bounds_hidden_urls_albums_and_multiple_matches():
    req = SearchRequest.create("@otherbot", "custom", date(2026, 10, 6), date(2026, 10, 6))
    mgr = manager({111: [post(6, at=datetime(2026, 10, 6, 19, tzinfo=timezone.utc)),
        post(5, "caption", at=datetime(2026, 10, 6, 18, 59, 59, tzinfo=timezone.utc),
             entities=[NS(url="https://t.me/OtherBot?start=1")], album=77),
        post(4, album=77), post(3, at=datetime(2026, 10, 5, 19, tzinfo=timezone.utc)),
        post(2, at=datetime(2026, 10, 5, 18, 59, 59, tzinfo=timezone.utc))]})
    report = await LiveSearchService(mgr).search([channel()], req)
    assert [ad.message_id for ad in report.channels[0].ads] == [5, 3]
    assert mgr.client.iter_messages.call_args.kwargs["limit"] == 501
    assert mgr.client.iter_messages.call_args.kwargs["offset_date"] > req.end


@pytest.mark.asyncio
async def test_not_found_is_distinct_from_failure_and_partial_matches_are_kept():
    mgr = manager({111: [post(1, "news")], 222: [ChannelPrivateError(request=None)],
                   333: [post(2), ConnectionError("network")], 444: [post(3)]})
    report = await LiveSearchService(mgr).search([channel(n) for n in (111, 222, 333, 444)], custom_request())
    assert [ch.status for ch in report.channels] == ["not_found", "error", "error", "found"]
    assert report.channels[2].ads[0].message_id == 2
    assert report.total_ads == 2


@pytest.mark.asyncio
async def test_flood_wait_stops_remaining_channels_and_prevents_early_retry():
    flood = FloodWaitError(request=None)
    flood.seconds = 300
    mgr = manager({111: [post(2), flood], 222: [post(3)]})
    service = LiveSearchService(mgr)
    report = await service.search([channel(), channel(222)], custom_request())
    assert [ch.status for ch in report.channels] == ["error", "error"]
    assert len(report.channels[0].ads) == 1
    assert mgr.client.iter_messages.call_count == 1
    with pytest.raises(SearchUnavailable, match="Telegram cheklovi"):
        await service.search([channel()], custom_request())


@pytest.mark.asyncio
async def test_search_reads_only_registered_active_channels_and_writes_no_ads(db_session, sample_channel, sample_ad):
    for tg_id, active, deleted in ((-100222, False, False), (-100333, True, True)):
        db_session.add(Channel(telegram_channel_id=tg_id, title="Excluded", url="https://t.me/excluded",
                               is_active=active, is_deleted=deleted))
    await db_session.flush()
    targets = await load_search_channels(db_session)
    assert [ch.id for ch in targets] == [sample_channel.id]
    mgr = manager({1987654321: [post(5)]})
    await LiveSearchService(mgr).search(targets, custom_request())
    assert await db_session.scalar(select(func.count()).select_from(Advertisement)) == 1
    assert await db_session.scalar(select(func.count()).select_from(ViewSnapshot)) == 0
    assert sample_ad.views == 15000
    assert sample_channel.last_scanned_at is None


@pytest.mark.asyncio
async def test_bulk_add_spaces_newlines_aliases_invalid_and_readd_by_id(db_session):
    def resolve(raw):
        if raw == "@broken":
            raise ChannelNotFoundError()
        tg_id = -100111 if "first" in raw.lower() or "renamed" in raw else -100222
        return dict(telegram_channel_id=tg_id, title=raw, username=raw.lstrip("@"), url="https://t.me/first")
    tg = NS(resolve_channel=AsyncMock(side_effect=resolve))
    service = ChannelService(db_session, tg)
    results = await service.add_channels("@first https://t.me/FIRST\n@broken\n@second @renamed")
    assert [item.status for item in results] == ["added", "duplicate", "error", "added", "duplicate"]
    all_channels = await service.list_all_channels()
    assert len(all_channels) == 2
    first = next(ch for ch in all_channels if ch.telegram_channel_id == -100111)
    await service.remove_channel(first.id)
    assert first.id not in [ch.id for ch in await service.list_active_channels()]
    result = (await service.add_channels("@renamed"))[0]
    assert result.status == "added"
    assert result.telegram_id == -100111
    assert len(await service.list_all_channels()) == 2
    assert (await service.channel_repo.get_by_tg_id(-100111)).id == first.id


@pytest.mark.asyncio
async def test_bulk_flood_saves_completed_items_without_hammering_rest(db_session):
    tg = NS(resolve_channel=AsyncMock(side_effect=[dict(telegram_channel_id=-100111, title="A", username="a", url="https://t.me/a"),
                                                  TelegramFloodWaitException(100)]))
    service = ChannelService(db_session, tg)
    result = await service.add_channels("@a @b @c")
    assert [item.status for item in result] == ["added", "error", "error"]
    assert tg.resolve_channel.await_count == 2
    assert len(await service.list_all_channels()) == 1


@pytest.mark.asyncio
async def test_channel_health_check_uses_id_instead_of_old_username(db_session, sample_channel):
    tg = NS(resolve_channel=AsyncMock(return_value=dict(title="New title", username="new_name", url="https://t.me/new_name",
                                                       telegram_channel_id=sample_channel.telegram_channel_id)))
    result = await ChannelService(db_session, tg).check_channel_health(sample_channel.id)
    assert result["status"] == "ok"
    tg.resolve_channel.assert_awaited_once_with(str(sample_channel.telegram_channel_id))


def example_report():
    first = channel()
    second = channel(222)
    return SearchReport(custom_request(), [
        ChannelSearchResult(first, [LiveAd(1, "https://t.me/first/1", datetime(2026, 10, 6, 10, tzinfo=timezone.utc), 60200, 7),
                                    LiveAd(2, "https://t.me/first/2", datetime(2026, 10, 12, 10, tzinfo=timezone.utc), 0, None)], "found", "Topildi"),
        ChannelSearchResult(second, [LiveAd(3, "https://t.me/second/3", datetime(2026, 10, 12, 10, tzinfo=timezone.utc), 2200, 0)], "found", "Topildi"),
        ChannelSearchResult(channel(333)),
        ChannelSearchResult(channel(444), status="error", detail="Kanalga kirish huquqi yo'q"),
    ], datetime(2026, 10, 22, tzinfo=timezone.utc))


def test_excel_matches_reference_matrix_and_keeps_unknown_separate_from_zero():
    report = example_report()
    report.channels[0].channel = SearchChannel(111, -100111, "=UNSAFE()", "https://t.me/first", "first", "umumiy", "@manager")
    wb = openpyxl.load_workbook(generate_live_excel(report))
    assert wb.sheetnames == ["REKLAMA HISOBOTI"]
    ws = wb.active
    assert [ws.cell(1, col).value for col in (1, 2, 3, 4)] == ["VILOYAT", "Kanal nomi", "Silkasi", "Silkasi (aloqa)"]
    assert "6 - oktyabr" in ws["E1"].value
    assert "12 - oktyabr" in ws["H1"].value
    assert ws["E2"].hyperlink.target == "https://t.me/first/1"
    assert ws["F2"].value == 60200 and ws["G2"].value == 7
    assert ws["E3"].fill.fgColor.rgb.endswith("FF0000")
    assert ws["H3"].value == "https://t.me/second/3"
    assert ws["I2"].value == 0 and ws["J2"].value is None
    assert ws["E4"].value == "Topilmadi"
    assert ws["E5"].fill.fgColor.rgb.endswith("FFEB9C")
    assert ws["B2"].value.startswith("'=")
    assert ws["D2"].hyperlink.target == "https://t.me/manager"
    assert ws.freeze_panes == "E2"
    assert ws.column_dimensions["D"].hidden is True


def test_empty_and_latest_excel_reports():
    report = example_report()
    report.channels = [ChannelSearchResult(channel())]
    ws = openpyxl.load_workbook(generate_live_excel(report)).active
    assert ws["E2"].value == "Topilmadi"
    report.request = SearchRequest.create("@otherbot", "latest")
    ws = openpyxl.load_workbook(generate_live_excel(report)).active
    assert ws["E1"].value == "Eng oxirgi reklama"
    assert ws["H1"].value == "Reklama sanasi"


@pytest.mark.asyncio
async def test_unauthorized_account_is_not_reported_as_no_ads():
    mgr = manager({})
    mgr.is_authorized.return_value = False
    with pytest.raises(SearchUnavailable, match="ulanmagan"):
        await LiveSearchService(mgr).search([channel()], custom_request())
    mgr.client.iter_messages.assert_not_called()


@pytest.mark.asyncio
async def test_mismatching_channel_id_never_searches_another_peer():
    mgr = manager({999: [post(1)]})
    mgr.client.get_entity.return_value = None
    mgr.client.get_entity.side_effect = lambda marked: NS(id=999, username="other", title="Other", broadcast=True)
    report = await LiveSearchService(mgr).search([channel()], custom_request())
    assert report.channels[0].status == "error"
    mgr.client.iter_messages.assert_not_called()

"""Post reports use Telegram counters, not history searches or persisted snapshots."""

import asyncio
import logging
from datetime import datetime, timezone
from types import SimpleNamespace as NS
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook
from telethon.errors import ChatAdminRequiredError, FloodWaitError, StatsMigrateError
from telethon.tl.functions.channels import GetFullChannelRequest
from telethon.tl.functions.stats import GetMessagePublicForwardsRequest
from telethon.tl.types import (
    Channel, ChatPhotoEmpty, Message, MessageEmpty, MessageReactions, MessageReplies,
    PeerChannel, PeerUser, PublicForwardMessage, PublicForwardStory, ReactionCount, ReactionEmoji,
)

from app.reports.post_excel import generate_post_excel
from app.services.post_report import (
    MISSING, PostReportService, PostReportUnavailable, metrics, parse_post_link,
)

NOW = datetime(2026, 10, 21, 3, 53, tzinfo=timezone.utc)
LINK = "https://t.me/imkonsavdomarkazi/5644"


def channel(identifier=100, name="imkonsavdomarkazi", group=False, members=17960):
    return Channel(id=identifier, title="IMKON" if identifier == 100 else f"Kanal {identifier}",
                   photo=ChatPhotoEmpty(), date=NOW, access_hash=identifier * 100,
                   username=name, broadcast=not group, megagroup=group, participants_count=members)


def post(identifier=5644, peer=100, views=700300, shares=228):
    return Message(id=identifier, peer_id=PeerChannel(peer), date=NOW, message="Post",
                   views=views, forwards=shares,
                   reactions=MessageReactions([ReactionCount(ReactionEmoji("❤"), 8)]),
                   replies=MessageReplies(replies=3, replies_pts=1))


class FakeClient:
    def __init__(self):
        self.session = NS(dc_id=2)
        self.source = channel()
        self.message = post()
        self.pages = {"": NS(count=0, forwards=[], chats=[], next_offset=None)}
        self.get_entity = AsyncMock(return_value=self.source)
        self.get_messages = AsyncMock(side_effect=lambda *a, **kw: self.message)
        self.rpc = AsyncMock(side_effect=self.respond)
        self._borrow_exported_sender = AsyncMock(return_value=NS(dc_id=4))
        self._return_exported_sender = AsyncMock()
        self._call = AsyncMock(side_effect=lambda sender, request, **kw: self.respond(request))

    def respond(self, request):
        if isinstance(request, GetFullChannelRequest):
            return NS(full_chat=NS(participants_count=17960, stats_dc=2))
        if isinstance(request, GetMessagePublicForwardsRequest):
            response = self.pages[request.offset]
            if isinstance(response, Exception):
                raise response
            return response
        raise AssertionError(type(request).__name__)

    async def __call__(self, request):
        return await self.rpc(request)


@pytest.fixture
def setup_service():
    client = FakeClient()
    manager = NS(client=client, is_authorized=AsyncMock(return_value=True))
    return PostReportService(manager), client, manager


@pytest.mark.parametrize("raw,expected", [
    (LINK, ("imkonsavdomarkazi", 5644)),
    (" t.me/s/imkonsavdomarkazi/5644?single ", ("imkonsavdomarkazi", 5644)),
    ("https://t.me/c/1234567/10", (-1000001234567, 10)),
    ("https://telegram.me/imkonsavdomarkazi/1", ("imkonsavdomarkazi", 1)),
])
def test_post_links(raw, expected):
    parsed = parse_post_link(raw)
    assert (parsed.channel, parsed.message_id) == expected


@pytest.mark.parametrize("raw", ["", "https://evil.example/imkon/5", "https://t.me@evil.example/imkon/5",
    "https://t.me/imkon", "https://t.me/imkon/0", "https://t.me/imkon/-1",
    "https://t.me/imkon/2147483648", "https://t.me/imkon/5/6", "https://t.me/+private",
    "https://t.me/imkon/1 https://t.me/imkon/2", "https://t.me/c/0/1",
    "https://t.me/c/10000000000/1", "https://t.me/c/１２３/1"])
def test_invalid_links(raw):
    with pytest.raises(ValueError, match="Post havolasini"):
        parse_post_link(raw)


@pytest.mark.asyncio
async def test_all_pages_telegram_totals_group_metrics_and_story_exclusion(setup_service):
    service, client, _ = setup_service
    first, group = channel(101, "first"), channel(102, "group", group=True)
    forwarded = PublicForwardMessage(post(10, 101, 73031, 12))
    group_post = PublicForwardMessage(post(11, 102, 28640, 5))
    # Public count includes a story; private count must use the unfiltered API total.
    story = PublicForwardStory(PeerUser(999), NS(id=1))
    client.pages = {
        "": NS(count=3, forwards=[forwarded, story], chats=[first], next_offset="page2"),
        "page2": NS(count=3, forwards=[forwarded, group_post], chats=[group], next_offset=None),
    }
    report = await service.collect(LINK)
    assert report.source.views == 700300
    assert report.source.members == 17960
    assert report.source.reactions == 8
    assert report.public_shares == 3 and report.private_shares == 225
    assert len(report.forwards) == 2 and report.forwards_complete
    assert report.forwards[0].views == 73031
    assert report.forwards[0].shares == 12
    assert report.forwards[0].replies == 3
    assert report.forwards[0].post_url == "https://t.me/first/10"
    assert report.forwards[1].kind == "Guruh"
    assert report.forwards[1].views == 28640
    stats = [call.args[0] for call in client.rpc.await_args_list
             if isinstance(call.args[0], GetMessagePublicForwardsRequest)]
    assert [request.offset for request in stats] == ["", "page2"]
    # There is no history iterator, database repository or search in this fake client.
    client.get_messages.assert_awaited_once()


@pytest.mark.asyncio
async def test_missing_public_stats_logs_info_not_error_and_does_not_invent_private_count(setup_service, caplog):
    service, client, _ = setup_service
    client.pages[""] = ChatAdminRequiredError(None)
    with caplog.at_level(logging.INFO, logger="imkon_ads_monitor"):
        report = await service.collect(LINK)
    assert report.source.views == 700300
    assert report.public_shares is None and report.private_shares is None
    assert report.forwards == [] and not report.forwards_complete
    assert MISSING in caplog.text
    assert all(record.levelno < logging.WARNING for record in caplog.records)


@pytest.mark.asyncio
async def test_unavailable_second_page_keeps_first_page_and_authoritative_total(setup_service):
    service, client, _ = setup_service
    client.pages = {
        "": NS(count=15, forwards=[PublicForwardMessage(post(10, 101))],
               chats=[channel(101, "first")], next_offset="p2"),
        "p2": ChatAdminRequiredError(None),
    }
    report = await service.collect(LINK)
    assert len(report.forwards) == 1 and not report.forwards_complete
    assert report.public_shares == 15 and report.private_shares == 213


@pytest.mark.asyncio
async def test_zero_and_missing_counts_remain_different(setup_service):
    service, client, _ = setup_service
    client.message = post(views=0, shares=0)
    client.message.reactions = None
    client.message.replies = None
    report = await service.collect(LINK)
    assert report.source.views == 0 and report.source.reactions is None
    assert report.public_shares == 0 and report.private_shares == 0
    client.message.forwards = None
    report = await service.collect(LINK)
    assert report.private_shares is None
    client.message.forwards = 0
    client.pages[""].count = 3
    report = await service.collect(LINK)
    assert report.private_shares is None  # Inconsistent counters do not become a fake zero.


@pytest.mark.asyncio
@pytest.mark.parametrize("missing", [None, MessageEmpty(5644, PeerChannel(100))])
async def test_missing_source_is_normal_information(setup_service, caplog, missing):
    service, client, _ = setup_service
    client.message = missing
    with caplog.at_level(logging.INFO, logger="imkon_ads_monitor"):
        with pytest.raises(PostReportUnavailable, match=MISSING):
            await service.collect(LINK)
    assert all(record.levelno < logging.WARNING for record in caplog.records)


@pytest.mark.asyncio
async def test_authorization_and_floodwait(setup_service):
    service, client, manager = setup_service
    manager.is_authorized.return_value = False
    with pytest.raises(PostReportUnavailable, match="avtorizatsiyadan"):
        await service.collect(LINK)
    client.get_entity.assert_not_awaited()
    manager.is_authorized.return_value = True
    client.pages[""] = FloodWaitError(None, capture=35)
    report = await service.collect(LINK)
    assert "35 soniya" in report.notice
    assert report.source.views == 700300 and not report.forwards_complete


@pytest.mark.asyncio
async def test_stats_dc_borrow_is_returned_on_cancellation(setup_service):
    service, client, _ = setup_service
    request = GetMessagePublicForwardsRequest(client.source, 5644, "", 100)
    client._call.side_effect = asyncio.CancelledError
    with pytest.raises(asyncio.CancelledError):
        await service._stats_request(request, 4)
    client._borrow_exported_sender.assert_awaited_once_with(4)
    client._return_exported_sender.assert_awaited_once()


@pytest.mark.asyncio
async def test_stats_migration_retries_on_exported_sender_without_migrating_main_client(setup_service):
    service, client, _ = setup_service
    client.rpc.side_effect = StatsMigrateError(None, capture=4)
    request = GetMessagePublicForwardsRequest(client.source, 5644, "", 100)
    result = await service._stats_request(request, None)
    assert result.count == 0
    client._borrow_exported_sender.assert_awaited_once_with(4)
    client._return_exported_sender.assert_awaited_once()
    assert client.session.dc_id == 2


@pytest.mark.asyncio
async def test_repeated_offset_stops_without_marking_complete(setup_service):
    service, client, _ = setup_service
    page = NS(count=0, forwards=[], chats=[], next_offset="same")
    client.pages = {"": page, "same": page}
    report = await service.collect(LINK)
    assert not report.forwards_complete
    assert client.rpc.await_count == 3  # full channel + two pages


@pytest.mark.asyncio
async def test_excel_only_two_tables_real_numbers_timezone_and_no_unsafe_formulas(setup_service):
    service, client, _ = setup_service
    client.source.title = '\x00=HYPERLINK("https://evil.example")'
    client.message.views = 0
    client.message.reactions = None
    destination = channel(101, "first")
    client.pages[""] = NS(count=1, forwards=[PublicForwardMessage(post(10, 101))],
                          chats=[destination], next_offset=None)
    report = await service.collect(LINK)
    workbook = load_workbook(generate_post_excel(report))
    assert workbook.sheetnames == ["Asosiy post", "Public forwardlar"]
    main, forward = workbook.worksheets
    assert main.max_row == 2 and main.max_column == 13
    assert forward.max_row == 2 and forward.max_column == 10
    assert main["G2"].value == 0
    assert main["H2"].value == MISSING
    assert main["K2"].value == 1 and main["L2"].value == 227
    assert main["F2"].value == datetime(2026, 10, 21, 8, 53)
    assert main["B2"].data_type == "s"
    assert main["E2"].hyperlink.target == LINK
    assert forward["E2"].hyperlink.target == "https://t.me/first/10"
    assert len(main.tables) == len(forward.tables) == 1
    assert not workbook.defined_names
    for sheet in workbook:
        assert not sheet.merged_cells and not sheet._charts and not sheet._images
        assert not any(cell.data_type == "f" for row in sheet for cell in row)


def test_missing_forward_counters_and_zero_reactions():
    message = post(views=None, shares=None)
    message.reactions = MessageReactions([])
    result = metrics(message, channel())
    assert result.views is None and result.shares is None and result.reactions == 0


@pytest.mark.asyncio
async def test_enrich_missing_forward_counters_from_exact_destination_post(setup_service):
    service, client, _ = setup_service
    destination = channel(101, "first", members=None)
    basic = post(10, 101, views=73031, shares=None)
    basic.reactions = basic.replies = None
    client.pages[""] = NS(count=1, forwards=[PublicForwardMessage(basic)],
                          chats=[destination], next_offset=None)
    complete = post(10, 101, views=73035, shares=12)
    client.get_messages.side_effect = lambda entity, ids: client.message if ids == 5644 else complete
    report = await service.collect(LINK)
    row = report.forwards[0]
    assert (row.views, row.shares, row.reactions, row.replies, row.members) == (73035, 12, 8, 3, 17960)
    assert row.post_url == "https://t.me/first/10"
    assert client.get_messages.await_count == 2
    assert client.get_messages.await_args_list[1].kwargs == {"ids": 10}


@pytest.mark.asyncio
async def test_unavailable_forward_enrichment_preserves_returned_views_without_warnings(setup_service, caplog):
    service, client, _ = setup_service
    basic = post(10, 101, views=73031, shares=None)
    basic.reactions = basic.replies = None
    client.pages[""] = NS(count=1, forwards=[PublicForwardMessage(basic)],
                          chats=[channel(101, "first")], next_offset=None)

    async def read(entity, ids):
        if ids == 5644:
            return client.message
        raise ChatAdminRequiredError(None)

    client.get_messages.side_effect = read
    with caplog.at_level(logging.INFO, logger="imkon_ads_monitor"):
        report = await service.collect(LINK)
    assert report.forwards[0].views == 73031
    assert report.forwards[0].shares is None
    assert report.forwards_complete
    assert MISSING in caplog.text
    assert all(record.levelno < logging.WARNING for record in caplog.records)


@pytest.mark.asyncio
async def test_real_connection_failure_is_distinguished_from_missing_data(setup_service, caplog):
    service, client, _ = setup_service
    client.pages[""] = ConnectionError("synthetic failure")
    with caplog.at_level(logging.INFO, logger="imkon_ads_monitor"):
        report = await service.collect(LINK)
    assert not report.forwards_complete and report.public_shares is None
    assert any(record.levelno == logging.WARNING and "ulanish tugamadi" in record.message
               for record in caplog.records)
    assert "synthetic failure" not in caplog.text
    assert "ulanish tugamadi" in report.notice
    assert not any(MISSING in record.message for record in caplog.records)


@pytest.mark.asyncio
async def test_unrelated_public_source_is_rejected_without_account_access(setup_service):
    service, client, manager = setup_service
    with pytest.raises(PostReportUnavailable, match="faqat.*imkonsavdomarkazi"):
        await service.collect("https://t.me/unrelated/55")
    manager.is_authorized.assert_not_awaited()
    client.get_entity.assert_not_awaited()
    client.get_messages.assert_not_awaited()


@pytest.mark.asyncio
async def test_private_link_cannot_probe_another_account_chat(setup_service):
    service, client, _ = setup_service
    with pytest.raises(PostReportUnavailable, match="faqat.*imkonsavdomarkazi"):
        await service.collect("https://t.me/c/1234567/10")
    client.get_entity.assert_awaited_once_with("imkonsavdomarkazi")
    client.get_messages.assert_not_awaited()
    client.rpc.assert_not_awaited()


@pytest.mark.asyncio
async def test_owned_source_private_style_link_matches_telegram_id(setup_service):
    service, client, _ = setup_service
    report = await service.collect("https://t.me/c/100/5644")
    assert report.source.post_url == LINK
    client.get_entity.assert_awaited_once_with("imkonsavdomarkazi")


@pytest.mark.asyncio
@pytest.mark.parametrize("failure", [ConnectionError("secret detail"),
    ValueError("Request was unsuccessful 5 time(s)")])
async def test_source_transport_failure_is_never_reported_as_missing(setup_service, failure):
    service, client, _ = setup_service
    client.get_messages.side_effect = failure
    with pytest.raises(PostReportUnavailable) as raised:
        await service.collect(LINK)
    assert MISSING not in str(raised.value)
    assert "secret detail" not in str(raised.value)


@pytest.mark.asyncio
async def test_rate_limit_in_enrichment_keeps_row_source_and_public_total(setup_service):
    service, client, _ = setup_service
    basic = post(10, 101, views=73031, shares=None)
    client.pages[""] = NS(count=15, forwards=[PublicForwardMessage(basic)],
                          chats=[channel(101, "first")], next_offset="more")

    async def read(entity, ids):
        if ids == 5644:
            return client.message
        raise FloodWaitError(None, capture=60)

    client.get_messages.side_effect = read
    report = await service.collect(LINK)
    assert len(report.forwards) == 1 and report.forwards[0].views == 73031
    assert report.public_shares == 15 and report.private_shares == 213
    assert report.source.views == 700300
    assert not report.forwards_complete and "60 soniya" in report.notice
    assert len([call for call in client.rpc.await_args_list
                if isinstance(call.args[0], GetMessagePublicForwardsRequest)]) == 1


@pytest.mark.asyncio
async def test_slow_api_response_returns_partial_report_and_cancels_request(setup_service):
    service, client, _ = setup_service
    service.request_timeout = 0.02
    stopped = asyncio.Event()
    respond = client.respond

    async def blocked(request):
        if not isinstance(request, GetMessagePublicForwardsRequest):
            return respond(request)
        try:
            await asyncio.Future()
        finally:
            stopped.set()

    client.rpc.side_effect = blocked
    report = await service.collect(LINK)
    assert stopped.is_set() and not report.forwards_complete
    assert report.source.views == 700300 and "ulanish tugamadi" in report.notice


@pytest.mark.asyncio
async def test_partial_excel_is_labelled_without_extra_status_tables(setup_service):
    service, client, _ = setup_service
    client.pages[""] = ChatAdminRequiredError(None)
    report = await service.collect(LINK)
    workbook = load_workbook(generate_post_excel(report))
    assert workbook.sheetnames == ["Asosiy post", "Public forwardlar (qisman)"]
    assert workbook.worksheets[0].max_row == 2
    assert workbook.worksheets[1].max_row == 1
    assert workbook.worksheets[0]["K2"].value == MISSING
    assert workbook.worksheets[0]["L2"].value == MISSING

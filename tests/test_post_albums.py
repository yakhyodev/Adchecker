"""An album is one logical post; media exposure counts must never be summed."""

from copy import deepcopy
from datetime import date, datetime
from threading import Event
from types import SimpleNamespace as NS
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook
from telethon.tl.types import MessageEmpty, PeerChannel, PublicForwardMessage
from telethon.errors import FloodWaitError

from app.core.timezone import TASHKENT_TZ
from app.reports.post_range_excel import generate_post_range_excel
from app.services.post_range import PostRange, RangeStore, decode
from app.services.post_report import PostReportService
from tests.test_post_range import RangeClient, dated, service
from tests.test_post_report import FakeClient, channel, post


def album_messages(first=10, peer=100, group=1234):
    result = []
    for offset, views in enumerate([6713, 7308, 7735, 8624, 7986]):
        message = post(first + offset, peer, views=views, shares=47)
        message.grouped_id = group
        message.date = datetime(2026, 10, 4, 9, 37, tzinfo=TASHKENT_TZ)
        result.append(message)
    return result


def reads(messages):
    def read(entity, ids):
        def one(identifier):
            return deepcopy(next((m for m in messages if m.id == identifier), MessageEmpty(identifier, PeerChannel(entity.id))))
        return [one(identifier) for identifier in ids] if isinstance(ids, list) else one(ids)
    return read


@pytest.fixture
def spool(tmp_path):
    store = RangeStore(tmp_path / "spool.sqlite3")
    yield store
    store.close()


@pytest.mark.asyncio
async def test_five_media_are_one_row_first_link_last_counter_and_one_stats_request(spool, tmp_path):
    client = RangeClient(album_messages())
    client.get_messages.side_effect = reads(client.history)
    await service(client).collect(PostRange(date(2026, 10, 4), date(2026, 10, 4)), spool, AsyncMock())
    assert spool.post_count == 1
    row = spool.db.execute("SELECT message_id,source FROM posts").fetchone()
    source = decode(row[1])
    assert row[0] == 10 and source.post_url.endswith("/10")
    assert source.views == 7986 and source.shares == 47
    assert source.views != sum(m.views for m in client.history) and source.views != 8624
    assert client.stats_calls == [(14, "")]
    client.get_messages.assert_awaited_once()
    assert len(client.get_messages.await_args.kwargs["ids"]) == 19
    path = tmp_path / "album.xlsx"
    generate_post_range_excel(spool, path, Event())
    wb = load_workbook(path)
    assert wb.sheetnames == ["Postlar", "Forwardlar"]
    assert wb["Postlar"].max_row == 2
    assert wb["Postlar"]["C2"].hyperlink.target.endswith("/10")
    assert wb["Postlar"]["D2"].value == 7986
    wb.close()


@pytest.mark.asyncio
async def test_album_media_lookup_waits_and_retries_without_advancing(spool):
    client = RangeClient(album_messages())
    read = reads(client.history)
    attempts = 0

    def limited(entity, ids):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise FloodWaitError(None, capture=3)
        return read(entity, ids)

    client.get_messages.side_effect = limited
    sleep = AsyncMock()
    await service(client, sleep).collect(PostRange(date(2026, 10, 4), date(2026, 10, 4)), spool, AsyncMock())
    assert attempts == 2 and spool.post_count == 1
    sleep.assert_awaited_once_with(4)
    assert client.stats_calls == [(14, "")]


@pytest.mark.asyncio
async def test_same_timestamp_without_group_id_remains_separate_posts(spool):
    client = RangeClient(album_messages())
    for item in client.history:
        item.grouped_id = None
    await service(client).collect(PostRange(date(2026, 10, 4), date(2026, 10, 4)), spool, AsyncMock())
    assert spool.post_count == 5


@pytest.mark.asyncio
async def test_history_retry_does_not_recollect_remaining_album_media(spool):
    client = RangeClient(album_messages())
    client.get_messages.side_effect = reads(client.history)
    client.fail_history_once = True
    await service(client).collect(PostRange(date(2026, 10, 4), date(2026, 10, 4)), spool, AsyncMock())
    assert spool.post_count == 1 and client.stats_calls == [(14, "")]
    assert [call["min_id"] for call in client.history_calls] == [0, 10]


@pytest.mark.asyncio
async def test_deleted_first_media_uses_first_surviving_link(spool):
    client = RangeClient(album_messages())
    client.get_messages.side_effect = reads(client.history[1:])
    await service(client).collect(PostRange(date(2026, 10, 4), date(2026, 10, 4)), spool, AsyncMock())
    assert spool.post_count == 1
    source = decode(spool.db.execute("SELECT source FROM posts").fetchone()[0])
    assert source.message_id == 11 and source.views == 7986


@pytest.mark.asyncio
async def test_album_crossing_start_boundary_is_not_misreported_as_new_post(spool):
    messages = album_messages()
    messages[0].date = datetime(2026, 10, 3, 23, 59, 59, tzinfo=TASHKENT_TZ)
    plain = dated(15, 4, 12)
    client = RangeClient(messages + [plain])
    client.get_messages.side_effect = reads(client.history)
    await service(client).collect(PostRange(date(2026, 10, 4), date(2026, 10, 4)), spool, AsyncMock())
    assert [row[0] for row in spool.db.execute("SELECT message_id FROM posts")] == [15]


@pytest.mark.asyncio
async def test_single_post_link_to_middle_media_returns_one_album():
    client = FakeClient()
    messages = album_messages()
    client.get_messages.side_effect = reads(messages)
    api = PostReportService(NS(client=client, is_authorized=AsyncMock(return_value=True)))
    report = await api.collect("https://t.me/imkonsavdomarkazi/12")
    assert report.source.message_id == 10 and report.source.views == 7986
    assert report.source.post_url.endswith("/10")
    assert report.source.shares == 47 and report.private_shares == 47


@pytest.mark.asyncio
@pytest.mark.parametrize("stream", [False, True])
async def test_forward_album_is_deduplicated_across_pages_and_other_peers_are_kept(spool, stream):
    client = FakeClient()
    first, second = album_messages(20, 101), album_messages(30, 201)
    all_messages = first + second
    client.get_messages.side_effect = reads(all_messages)
    client.pages = {
        "": NS(count=10, forwards=[PublicForwardMessage(m) for m in first[2:] + second],
               chats=[channel(101, "first"), channel(201, "second")], next_offset="more"),
        "more": NS(count=10, forwards=[PublicForwardMessage(m) for m in first[:2]],
                   chats=[channel(101, "first")], next_offset=None),
    }
    api = PostReportService(NS(client=client))
    report = await api.collect_message(client.source, client.message, spool.sink(5644) if stream else None)
    if stream:
        rows = [decode(row[0]) for row in spool.db.execute("SELECT payload FROM forwards ORDER BY message_id")]
    else:
        rows = report.forwards
    assert len(rows) == report.forward_count == 2
    assert [row.message_id for row in rows] == [20, 30]
    assert [row.views for row in rows] == [7986, 7986]
    assert report.public_shares == 10  # Telegram's authoritative total is preserved.


@pytest.mark.asyncio
async def test_missing_album_response_does_not_fabricate_views_or_fall_back_to_maximum(spool):
    client = RangeClient(album_messages())
    client.get_messages.return_value = None
    client.get_messages.side_effect = None
    await service(client).collect(PostRange(date(2026, 10, 4), date(2026, 10, 4)), spool, AsyncMock())
    assert spool.post_count == spool.partial_count == 1
    source = decode(spool.db.execute("SELECT source FROM posts").fetchone()[0])
    assert source.views is source.shares is None

"""The supplied Top Reklama layout: one channel per row, dates across columns."""

import io
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.security import sanitize_excel_cell
from app.core.timezone import format_tashkent, to_tashkent
from app.reports.formatters import GRID_BORDER
from app.services.live_search import SearchReport

MISSING_FILL = PatternFill("solid", fgColor="FF0000")
ERROR_FILL = PatternFill("solid", fgColor="FFEB9C")
COMPACT_COUNT = '[>=1000000]0.0,,"m";[>=1000]0.0,"k";0'
MONTHS = ("", "yanvar", "fevral", "mart", "aprel", "may", "iyun", "iyul",
          "avgust", "sentyabr", "oktyabr", "noyabr", "dekabr")


def generate_live_excel(report: SearchReport) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    latest = report.request.mode == "latest"
    grouped = []
    capacities = Counter()
    for result in report.channels:
        by_day = defaultdict(list)
        for ad in sorted(result.ads, key=lambda ad: (ad.posted_at, ad.message_id)):
            by_day[to_tashkent(ad.posted_at).date()].append(ad)
        grouped.append(by_day)
        for day, ads in by_day.items():
            capacities[day] = max(capacities[day], len(ads))
    slots = [(day, n) for day in sorted(capacities) for n in range(capacities[day])]
    if latest or not slots:
        slots = [(None, 0)]

    # Keep every result if an unusually large request exceeds Excel's column limit.
    per_sheet = (16384 - 6) // 3
    for offset in range(0, len(slots), per_sheet):
        section = slots[offset:offset + per_sheet]
        ws = wb.create_sheet("REKLAMA HISOBOTI" if offset == 0 else f"DAVOMI {offset // per_sheet + 1}")
        headers = ["VILOYAT", "Kanal nomi", "Silkasi", "Silkasi (aloqa)"]
        for number, (day, slot) in enumerate(section, offset + 1):
            label = "Eng oxirgi reklama" if latest else (
                f"{number} | {day.day} - {MONTHS[day.month]} {day.year}" if day else "Tanlangan davr")
            headers.extend([label, "Ko'rilishlar", "Ulashishlar"])
        if latest:
            headers.append("Reklama sanasi")
        headers.append("Holat")
        ws.append(headers)
        for row_index, (result, by_day) in enumerate(zip(report.channels, grouped), 2):
            ch = result.channel
            row = [sanitize_excel_cell(ch.region), sanitize_excel_cell(ch.title), sanitize_excel_cell(ch.url),
                   sanitize_excel_cell(ch.contact)]
            missing_columns = []
            for day, slot in section:
                ads = result.ads if latest else by_day.get(day, [])
                ad = ads[slot] if slot < len(ads) else None
                if ad:
                    row.extend([ad.url, ad.views, ad.forwards])
                else:
                    missing_columns.append(len(row) + 1)
                    row.extend(["Topilmadi" if result.status != "error" else result.detail, None, None])
            if latest:
                row.append(to_tashkent(result.ads[0].posted_at).replace(tzinfo=None) if result.ads else None)
            row.append(result.detail)
            ws.append(row)
            for column in missing_columns:
                ws.cell(row_index, column).fill = ERROR_FILL if result.status == "error" else MISSING_FILL
            if result.status == "error":
                ws.cell(row_index, len(headers)).fill = ERROR_FILL

        for row in ws:
            for cell in row:
                cell.border = GRID_BORDER
                cell.alignment = Alignment(horizontal="center" if cell.row == 1 else "left",
                                           vertical="center", wrap_text=True)
                cell.font = Font(name="Calibri", size=14 if cell.row == 1 else 11, bold=cell.row == 1)
                if cell.row > 1 and cell.column in (1, 2):
                    cell.font = Font(name="Calibri", size=14)
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if cell.row > 1 and isinstance(cell.value, str) and cell.value.startswith("https://t.me/"):
                    cell.hyperlink = cell.value
                    cell.font = Font(name="Calibri", size=11, color="0563C1", underline="single")
                elif cell.row > 1 and cell.column == 4 and isinstance(cell.value, str) and cell.value.startswith("'@"):
                    cell.hyperlink = "https://t.me/" + cell.value[2:]
                elif isinstance(cell.value, int):
                    cell.number_format = COMPACT_COUNT
                elif hasattr(cell.value, "strftime"):
                    cell.number_format = "dd.mm.yyyy hh:mm"
        widths = [13, 23, 42, 44]
        for _ in section:
            widths.extend([49, 16, 16])
        if latest:
            widths.append(23)
        widths.append(40)
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        # The reference keeps the advertising contact column available but hidden.
        ws.column_dimensions["D"].hidden = True
        ws.row_dimensions[1].height = 42
        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = 34
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
        ws.append([])
        ws.append(["Kalit so'zlar", sanitize_excel_cell(", ".join(report.request.keywords))])
        ws.append(["Davr", report.request.label])
        ws.append(["Tekshirildi", format_tashkent(report.checked_at) + " (Toshkent)"])
        ws.append(["Belgilar", "Qizil: topilmadi. Sariq: tekshiruv tugamadi. Bo'sh son: Telegram bermagan."])
        footer_start = len(report.channels) + 3
        for row in range(footer_start, ws.max_row + 1):
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=min(7, len(headers)))
            ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[row].height = 30
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

"""Plain tables for the one-shot post report, without charts or footer notes."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.core.security import sanitize_excel_cell
from app.core.timezone import to_tashkent
from app.services.post_report import MISSING, PostReport
from app.reports.xlsx_validation import validate_xlsx

HEADERS = ["Turi", "Kanal / guruh nomi", "Kanal / guruh havolasi", "Obunachilar / a'zolar",
           "Post havolasi", "Chiqarilgan vaqt (Toshkent)", "Ko'rishlar", "Reaksiyalar",
           "Ulashishlar", "Izohlar"]


def _safe(value):
    if value is None:
        return MISSING
    if isinstance(value, str):
        value = "".join(ch for ch in value if ch in "\n\t\r" or
                        0x20 <= ord(ch) <= 0xD7FF or 0xE000 <= ord(ch) <= 0xFFFD or
                        0x10000 <= ord(ch) <= 0x10FFFF)
    return sanitize_excel_cell(value)


def _row(post):
    return [_safe(value) for value in [post.kind, post.title, post.channel_url, post.members, post.post_url,
            to_tashkent(post.published_at).replace(tzinfo=None) if post.published_at else None,
            post.views, post.reactions, post.shares, post.replies]]


def generate_post_excel(report: PostReport) -> BytesIO:
    wb = Workbook()
    source = wb.active
    source.title = "Asosiy post"
    source.append(HEADERS + ["Public Shares", "Private Shares", "Hisobot vaqti (Toshkent)"])
    source.append(_row(report.source) + [_safe(report.public_shares), _safe(report.private_shares),
                  to_tashkent(report.checked_at).replace(tzinfo=None)])
    forwards = wb.create_sheet("Public forwardlar" if report.forwards_complete else "Public forwardlar (qisman)")
    forwards.append(HEADERS)
    for post in report.forwards:
        forwards.append(_row(post))
    widths = [14, 34, 40, 24, 49, 29, 23, 23, 23, 23, 23, 23, 29]
    for number, ws in enumerate(wb, 1):
        for row in ws:
            for cell in row:
                # Even unusual external strings must be stored as text, never as formulas.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11, bold=cell.row == 1)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if isinstance(cell.value, int):
                    cell.number_format = "#,##0"
                elif hasattr(cell.value, "strftime"):
                    cell.number_format = "dd.mm.yyyy hh:mm"
                elif cell.row > 1 and isinstance(cell.value, str) and cell.value.startswith("https://t.me/"):
                    cell.hyperlink = cell.value
                    cell.font = Font(name="Calibri", size=11, color="0563C1", underline="single")
        for column in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(column)].width = widths[column - 1]
        ws.row_dimensions[1].height = 34
        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = 38
        # Excel rejects a header-only ListObject even though openpyxl reads it.
        # Keep the styled header, without inventing an empty data row.
        if ws.max_row > 1:
            table = Table(displayName=f"PostReport{number}", ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        else:
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="24476A")
                cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        ws.freeze_panes = "C2"
        ws.sheet_view.showGridLines = False
        ws.print_title_rows = "1:1"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    buffer = BytesIO()
    wb.save(buffer)
    validate_xlsx(buffer)
    buffer.seek(0)
    return buffer

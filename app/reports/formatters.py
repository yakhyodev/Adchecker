from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# IMKON Brand Colors
PRIMARY_COLOR = "1E3A8A"  # Deep Navy Blue
ACCENT_COLOR = "3B82F6"   # Electric Blue
ZEBRA_COLOR = "F1F5F9"    # Light Slate Gray
HEADER_TEXT_COLOR = "FFFFFF"

# Fonts
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=HEADER_TEXT_COLOR)
BOLD_FONT = Font(name="Calibri", size=11, bold=True)
REGULAR_FONT = Font(name="Calibri", size=11)
LINK_FONT = Font(name="Calibri", size=11, color="2563EB", underline="single")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1E3A8A")

# Fills
HEADER_FILL = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
ZEBRA_FILL = PatternFill(start_color=ZEBRA_COLOR, end_color=ZEBRA_COLOR, fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

# Alignments
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

# Borders
THIN_SIDE = Side(border_style="thin", color="CBD5E1")
GRID_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

# Number Formats
NUMBER_FORMAT = "#,##0"
PERCENT_FORMAT = "0.0%"

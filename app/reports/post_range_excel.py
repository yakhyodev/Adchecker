"""Compact row-based reports streamed from the existing SQLite spool."""

from datetime import date, datetime
from itertools import islice
from pathlib import Path
from threading import Event

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.timezone import get_tashkent_now, to_tashkent
from app.reports.post_excel import _safe
from app.reports.xlsx_validation import validate_xlsx
from app.services.post_range import decode

SOURCE = ["№", "Chiqarilgan sana va vaqt", "Post havolasi", "Ko'rishlar",
          "Ulashishlar", "Public Shares", "Private Shares"]
FORWARD = ["№", "Asosiy post №", "Kanal / guruh nomi", "Kanal / guruh havolasi",
           "Forward havolasi", "Chiqarilgan sana va vaqt", "Ko'rishlar", "Ulashishlar"]
MISSING = "Ma'lumot mavjud emas"
# At most two links per forward: below Excel's 65,530 links per sheet.
# Bound openpyxl's retained relationship objects as well as worksheet rows.
MAX_DATA_ROWS = 20000
HEADER_FILL = PatternFill("solid", fgColor="153047")
DAY_FILL = PatternFill("solid", fgColor="EAF3F8")
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=11, color="153047")
LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")
DAY_BORDER = Border(top=Side(style="thin", color="A8C4D8"))
RANKED = """WITH ranked AS (
    SELECT *, ROW_NUMBER() OVER (ORDER BY day,json_extract(source,'$.published_at'),message_id) AS number
    FROM posts)
"""


class RangeExportCancelled(Exception):
    pass


class DiskWorksheetWriter(WorksheetWriter):
    def cleanup(self):
        Path(self.out).unlink(missing_ok=True)


def _local(value):
    return to_tashkent(value).replace(tzinfo=None) if value else None


def post_rows(store, post_offset=0, post_limit=-1):
    return store.db.execute(RANKED + "SELECT * FROM ranked ORDER BY number LIMIT ? OFFSET ?",
                            (post_limit, post_offset))


def forward_rows(store, post_offset=0, post_limit=-1):
    # Global source ordinals survive splitting and link back to Postlar.
    return store.db.execute(RANKED + """
        SELECT p.number,p.day,f.payload FROM forwards f
        JOIN ranked p ON p.message_id=f.source_id
        WHERE p.number>? AND (?<0 OR p.number<=?)
        ORDER BY p.number,f.published,f.destination_id,f.message_id
        """, (post_offset, post_limit, post_offset + post_limit))


def forward_count(store, post_offset=0, post_limit=-1):
    return store.db.execute(RANKED + """SELECT COUNT(*) FROM forwards f
        JOIN ranked p ON p.message_id=f.source_id
        WHERE p.number>? AND (?<0 OR p.number<=?)""",
        (post_offset, post_limit, post_offset + post_limit)).fetchone()[0]


def _cells(ws, values, header=False, *, blue=False, boundary=False):
    for value in values:
        value = MISSING if value is None else _safe(value)
        cell = WriteOnlyCell(ws, value=value)
        if isinstance(value, str):
            cell.data_type = "s"
        cell.font = HEADER_FONT if header else BODY_FONT
        cell.fill = HEADER_FILL if header else DAY_FILL if blue else WHITE_FILL
        cell.alignment = Alignment(horizontal="center" if header else "right" if isinstance(value, (int, float)) else "left",
                                   vertical="center", wrap_text=True)
        if boundary and not header:
            cell.border = DAY_BORDER
        if isinstance(value, (int, float)):
            cell.number_format = "#,##0"
        elif isinstance(value, datetime):
            cell.number_format = "dd.mm.yyyy hh:mm"
        elif not header and isinstance(value, str) and value.startswith("https://t.me/"):
            cell.hyperlink = value
            cell.font = LINK_FONT
        yield cell


def _sheet(wb, title, headers, widths, path, *, detail=True):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "153047"
    if detail:
        ws.freeze_panes = "C2"
        ws.print_title_rows = "1:1"
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.row_dimensions[1].height = 42
    ws.sheet_format.defaultRowHeight = 48
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = Worksheet.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws._writer = DiskWorksheetWriter(ws, str(Path(path).with_suffix(f".sheet{len(wb.worksheets)}.xml")))
    ws._writer.write_top()
    ws.append(_cells(ws, headers, header=True))
    return ws


def _summary(store, stop):
    count, total, best, latest = 0, 0, None, None
    all_known = True
    for row in post_rows(store):
        if stop.is_set():
            raise RangeExportCancelled()
        post = decode(row[3])
        count += 1
        all_known &= post.views is not None
        total += post.views or 0
        if post.views is not None and (best is None or post.views > best[0]):
            best = post.views, count, post.post_url
        checked = datetime.fromisoformat(row[6])
        latest = checked if latest is None or checked > latest else latest
    nforwards, views, known, best_channel = 0, 0, True, None
    for title, amount, observed, n in store.db.execute("""SELECT d.title,
            SUM(json_extract(f.payload,'$.views')),COUNT(json_extract(f.payload,'$.views')),COUNT(*)
            FROM forwards f JOIN destinations d ON d.id=f.destination_id
            GROUP BY d.id ORDER BY SUM(json_extract(f.payload,'$.views')) DESC,d.id"""):
        if stop.is_set():
            raise RangeExportCancelled()
        nforwards += n
        views += amount or 0
        known &= observed == n
        if best_channel is None and observed == n:
            best_channel = title
    if not known:
        best_channel = None  # Missing counters can change the ranking.
    days = store.db.execute("SELECT MIN(day),MAX(day) FROM posts").fetchone()
    selected = getattr(store, "selected_range", None)
    start, end = (selected.start, selected.end) if selected else days
    empty_complete = not nforwards and not store.partial_count and store.collection_complete
    return [
        ("Hisobot davri", f"{start} — {end}"),
        ("Yaratilgan vaqt (Toshkent)", get_tashkent_now().replace(tzinfo=None)),
        ("Statistika olingan vaqt (Toshkent)", _local(latest)),
        ("Holat", "Qisman" if store.partial_count or not store.collection_complete else "To'liq"),
        ("Asosiy postlar soni", count),
        ("Asosiy postlarning umumiy ko'rishlari", total if count and all_known else None),
        ("Aniqlangan forwardlar soni", nforwards),
        ("Forwardlarning umumiy ko'rishlari", views if nforwards and known else 0 if empty_complete else None),
        ("Forward qilingan kanal va guruhlar soni", store.db.execute("SELECT COUNT(DISTINCT destination_id) FROM forwards").fetchone()[0]),
        ("Eng ko'p ko'rilgan asosiy post №", best[1] if best and all_known else None),
        ("Eng ko'p ko'rilgan asosiy post", best[2] if best and all_known else None),
        ("Eng ko'p forward ko'rishlari to'plagan kanal", best_channel),
    ]


def generate_post_range_excel(store, path, stop: Event, *, post_offset=0, post_limit=-1,
                             forward_offset=0, forward_limit=-1):
    """One row per record. Native links and rows are bounded per workbook."""
    count = store.post_count - post_offset if post_limit < 0 else post_limit
    nforwards = forward_count(store, post_offset, post_limit) - forward_offset if forward_limit < 0 else forward_limit
    if max(count, nforwards) > MAX_DATA_ROWS:
        raise ValueError("Excel faylini qatorlar bo'yicha qismlarga ajrating.")
    wb = Workbook(write_only=True)
    try:
        summary = _sheet(wb, "Umumiy xulosa", ["IMKON postlar hisoboti", "Qiymat"], [49, 65], path, detail=False)
        for label, value in _summary(store, stop):
            summary.append(_cells(summary, [label, value]))
        if count != store.post_count or nforwards != forward_count(store):
            for label, value in (("Fayldagi postlar soni", count), ("Fayldagi forwardlar soni", nforwards)):
                summary.append(_cells(summary, [label, value]))
        posts = _sheet(wb, "Postlar", SOURCE, [8, 25, 49, 23, 23, 23, 23], path)
        first_day = store.db.execute("SELECT MIN(day) FROM posts").fetchone()[0]
        first_day = date.fromisoformat(first_day) if first_day else date(2000, 1, 1)
        previous, rows = None, 1
        for row in post_rows(store, post_offset, post_limit):
            if stop.is_set():
                raise RangeExportCancelled()
            _, _, day, data, public, private, _, _, number = row
            post = decode(data)
            blue = (date.fromisoformat(day) - first_day).days % 2 == 0
            posts.append(_cells(posts, [number, _local(post.published_at), post.post_url, post.views,
                                       post.shares, public, private], blue=blue, boundary=day != previous))
            previous, rows = day, rows + 1
        if rows > 1:
            posts.auto_filter.ref = f"A1:G{rows}"
        forwards = _sheet(wb, "Forwardlar", FORWARD, [8, 16, 45, 45, 49, 25, 23, 23], path)
        previous, rows = None, 1
        end = None if forward_limit < 0 else forward_offset + forward_limit
        for index, (number, day, payload) in enumerate(islice(forward_rows(store, post_offset, post_limit), forward_offset, end), forward_offset + 1):
            if stop.is_set():
                raise RangeExportCancelled()
            post = decode(payload)
            blue = (date.fromisoformat(day) - first_day).days % 2 == 0
            forwards.append(_cells(forwards, [index, number, post.title, post.channel_url, post.post_url,
                                              _local(post.published_at), post.views, post.shares],
                                   blue=blue, boundary=day != previous))
            previous, rows = day, rows + 1
        if rows > 1:
            forwards.auto_filter.ref = f"A1:H{rows}"
        if stop.is_set():
            raise RangeExportCancelled()
        wb.save(path)
        validate_xlsx(path)
    finally:
        wb.close()
        for ws in wb.worksheets:
            writer = getattr(ws, "_writer", None)
            if writer is not None:
                try:
                    if not ws.closed:
                        ws.close()
                finally:
                    writer.cleanup()
    if stop.is_set():
        raise RangeExportCancelled()

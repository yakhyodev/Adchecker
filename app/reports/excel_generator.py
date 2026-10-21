import io
from datetime import datetime, timezone
from typing import Dict, Any, List
import openpyxl
from openpyxl.utils import get_column_letter

from app.core.timezone import to_tashkent, format_tashkent, get_tashkent_now
from app.core.security import sanitize_excel_cell
from app.reports.formatters import (
    HEADER_FONT,
    BOLD_FONT,
    REGULAR_FONT,
    LINK_FONT,
    TITLE_FONT,
    HEADER_FILL,
    ZEBRA_FILL,
    SUBHEADER_FILL,
    ALIGN_LEFT,
    ALIGN_CENTER,
    ALIGN_RIGHT,
    GRID_BORDER,
    NUMBER_FORMAT,
)


def generate_ads_excel_report(report_data: Dict[str, Any]) -> io.BytesIO:
    """Generates a professional 4-sheet Excel report from report_data."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # 1. Sheet 1: REKLAMALAR
    _create_ads_sheet(wb, report_data.get("ads", []))

    # 2. Sheet 2: KANALLAR STATISTIKASI
    _create_channels_sheet(wb, report_data.get("channel_stats", []))

    # 3. Sheet 3: UMUMIY HISOBOT
    _create_summary_sheet(wb, report_data)

    # 4. Sheet 4: KUNLIK STATISTIKA
    _create_daily_sheet(wb, report_data.get("daily_stats", []))

    # Save to memory buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _create_ads_sheet(wb: openpyxl.Workbook, ads: List[Any]) -> None:
    ws = wb.create_sheet(title="REKLAMALAR")
    headers = [
        "№",
        "Kanal nomi",
        "Kanal havolasi",
        "Reklama joylashtirilgan sana",
        "Reklama joylashtirilgan vaqt",
        "Reklama posti havolasi",
        "Ko'rishlar",
        "Reaksiyalar",
        "Forwardlar",
        "Media turi",
        "Statistika yangilangan vaqt",
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = GRID_BORDER
    ws.row_dimensions[1].height = 28

    for idx, ad in enumerate(ads, start=1):
        ch = ad.channel
        ch_title = ch.title if ch else "Noma'lum"
        ch_url = ch.url if ch else ""

        local_posted = to_tashkent(ad.posted_at)
        p_date = local_posted.strftime("%Y-%m-%d") if local_posted else ""
        p_time = local_posted.strftime("%H:%M") if local_posted else ""
        u_time = format_tashkent(ad.last_updated_at)

        row_data = [
            idx,
            sanitize_excel_cell(ch_title),
            sanitize_excel_cell(ch_url),
            p_date,
            p_time,
            sanitize_excel_cell(ad.post_url),
            ad.views if ad.views is not None else "",
            ad.reactions if ad.reactions is not None else "",
            ad.forwards if ad.forwards is not None else "",
            sanitize_excel_cell(ad.media_type or "text"),
            u_time,
        ]
        ws.append(row_data)

        row_num = idx + 1
        fill_to_apply = ZEBRA_FILL if idx % 2 == 0 else None

        for c_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_num, column=c_idx)
            cell.font = REGULAR_FONT
            cell.border = GRID_BORDER
            if fill_to_apply:
                cell.fill = fill_to_apply

            # Specific column alignments & styles
            if c_idx in (1, 4, 5, 10, 11):
                cell.alignment = ALIGN_CENTER
            elif c_idx in (7, 8, 9):
                cell.alignment = ALIGN_RIGHT
                if isinstance(cell.value, (int, float)):
                    cell.number_format = NUMBER_FORMAT
            else:
                cell.alignment = ALIGN_LEFT

            # Clickable hyperlink for channel and post url
            if c_idx == 3 and ch_url and ch_url.startswith("http"):
                cell.hyperlink = ch_url
                cell.font = LINK_FONT
            elif c_idx == 6 and ad.post_url and ad.post_url.startswith("http"):
                cell.hyperlink = ad.post_url
                cell.font = LINK_FONT

        ws.row_dimensions[row_num].height = 20

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(ads) + 1}"
    _auto_fit_columns(ws)


def _create_channels_sheet(wb: openpyxl.Workbook, channel_stats: List[Dict[str, Any]]) -> None:
    ws = wb.create_sheet(title="KANALLAR STATISTIKASI")
    headers = [
        "Kanal nomi",
        "Kanal havolasi",
        "Reklama soni",
        "Umumiy ko'rishlar",
        "O'rtacha ko'rishlar",
        "Maksimal ko'rishlar",
        "Reaksiyalar yig'indisi",
        "Oxirgi reklama sanasi",
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = GRID_BORDER
    ws.row_dimensions[1].height = 28

    for idx, ch in enumerate(channel_stats, start=1):
        row_data = [
            sanitize_excel_cell(ch["title"]),
            sanitize_excel_cell(ch["url"]),
            ch["ads_count"],
            ch["total_views"],
            ch["avg_views"],
            ch["max_views"],
            ch["total_reactions"],
            ch["last_ad_date"],
        ]
        ws.append(row_data)

        row_num = idx + 1
        fill_to_apply = ZEBRA_FILL if idx % 2 == 0 else None

        for c_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_num, column=c_idx)
            cell.font = REGULAR_FONT
            cell.border = GRID_BORDER
            if fill_to_apply:
                cell.fill = fill_to_apply

            if c_idx in (3, 4, 5, 6, 7):
                cell.alignment = ALIGN_RIGHT
                cell.number_format = NUMBER_FORMAT
            elif c_idx == 8:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

            if c_idx == 2 and ch["url"] and ch["url"].startswith("http"):
                cell.hyperlink = ch["url"]
                cell.font = LINK_FONT

        ws.row_dimensions[row_num].height = 20

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(channel_stats) + 1}"
    _auto_fit_columns(ws)


def _create_summary_sheet(wb: openpyxl.Workbook, report_data: Dict[str, Any]) -> None:
    ws = wb.create_sheet(title="UMUMIY HISOBOT")

    # Title block
    ws["B2"] = "IMKON SAVDO MARKAZI — REKLAMA MONITORING HISOBOTI"
    ws["B2"].font = TITLE_FONT
    ws.row_dimensions[2].height = 25

    metrics = [
        ("Hisobot davri", f"{report_data.get('start_dt', '')} — {report_data.get('end_dt', '')}"),
        ("Reklamalar soni", report_data.get("total_ads", 0)),
        ("Kanallar soni", report_data.get("total_channels", 0)),
        ("Ko'rishlar yig'indisi", report_data.get("total_views", 0)),
        ("O'rtacha ko'rish", report_data.get("avg_views", 0)),
        ("Reaksiyalar yig'indisi", report_data.get("total_reactions", 0)),
        ("Forwardlar yig'indisi", report_data.get("total_forwards", 0)),
        ("Hisobot shakllantirilgan vaqt", format_tashkent(get_tashkent_now())),
        ("Statistikalar oxirgi yangilangan vaqt", report_data.get("last_updated_at", "")),
    ]

    # Add Top Ad info if available
    top_ad = report_data.get("top_ad")
    if top_ad:
        metrics.append(
            ("Eng ko'p ko'rilgan reklama", f"{top_ad['channel_title']} ({top_ad['views']:,} ko'rish)")
        )

    start_row = 4
    for idx, (label, val) in enumerate(metrics):
        r = start_row + idx
        ws[f"B{r}"] = label
        ws[f"C{r}"] = val

        c1 = ws[f"B{r}"]
        c2 = ws[f"C{r}"]

        c1.font = BOLD_FONT
        c1.fill = SUBHEADER_FILL
        c1.border = GRID_BORDER
        c1.alignment = ALIGN_LEFT

        c2.font = REGULAR_FONT
        c2.border = GRID_BORDER

        if isinstance(val, (int, float)):
            c2.alignment = ALIGN_RIGHT
            c2.number_format = NUMBER_FORMAT
        else:
            c2.alignment = ALIGN_LEFT

        ws.row_dimensions[r].height = 22

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 40


def _create_daily_sheet(wb: openpyxl.Workbook, daily_stats: List[Dict[str, Any]]) -> None:
    ws = wb.create_sheet(title="KUNLIK STATISTIKA")
    headers = [
        "Sana",
        "Reklamalar soni",
        "Ko'rishlar yig'indisi",
        "O'rtacha ko'rishlar",
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = GRID_BORDER
    ws.row_dimensions[1].height = 28

    for idx, d in enumerate(daily_stats, start=1):
        row_data = [
            d["date"],
            d["ads_count"],
            d["total_views"],
            d["avg_views"],
        ]
        ws.append(row_data)

        row_num = idx + 1
        fill_to_apply = ZEBRA_FILL if idx % 2 == 0 else None

        for c_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_num, column=c_idx)
            cell.font = REGULAR_FONT
            cell.border = GRID_BORDER
            if fill_to_apply:
                cell.fill = fill_to_apply

            if c_idx == 1:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_RIGHT
                cell.number_format = NUMBER_FORMAT

        ws.row_dimensions[row_num].height = 20

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(daily_stats) + 1}"
    _auto_fit_columns(ws)


def _auto_fit_columns(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    """Adjusts worksheet column widths automatically."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
